import os, sys, subprocess, json, re
from pathlib import Path

sys.path.insert(0, '/Users/denis/Work/ai-files_reader/markitdown-ai/src')
from dotenv import load_dotenv
load_dotenv('/Users/denis/Work/ai-files_reader/markitdown-ai/.env')

LAB_DIR = '/Users/denis/Work/WinForms/лабы_про_телефонный_справочник'
OUT_DIR = '/tmp/lab_extract'
IMG_DIR = os.path.join(OUT_DIR, 'images')
os.makedirs(IMG_DIR, exist_ok=True)

# --- Extract text ---

def extract_text_doc(path):
    import textract
    return textract.process(path).decode('utf-8')

def extract_text_docx(path):
    from docx import Document
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs]
    return '\n'.join(paragraphs)

results = {}
for fname in os.listdir(LAB_DIR):
    fpath = os.path.join(LAB_DIR, fname)
    if fname.lower().endswith('.doc') and not fname.lower().endswith('.docx'):
        text = extract_text_doc(fpath)
        results[fname] = text
        with open(os.path.join(OUT_DIR, fname + '.txt'), 'w', encoding='utf-8') as f:
            f.write(text)
    elif fname.lower().endswith('.docx'):
        text = extract_text_docx(fpath)
        results[fname] = text
        with open(os.path.join(OUT_DIR, fname + '.txt'), 'w', encoding='utf-8') as f:
            f.write(text)

# --- Extract images ---

def extract_images_doc(path, out_dir):
    import olefile
    ole = olefile.OleFileIO(path)
    data = ole.openstream('Data').read()
    # Find PNG signatures
    idx = 0
    count = 0
    while True:
        sig = data.find(b'\x89PNG\r\n\x1a\n', idx)
        if sig == -1:
            break
        end = data.find(b'IEND\xaeB`\x82', sig)
        if end == -1:
            break
        end += 8
        out_path = os.path.join(out_dir, f'{Path(path).stem}_{count:02d}.png')
        with open(out_path, 'wb') as f:
            f.write(data[sig:end])
        count += 1
        idx = end
    print(f'Extracted {count} PNG from {path}')

def extract_images_docx(path, out_dir):
    from docx import Document
    doc = Document(path)
    count = 0
    for rel in doc.part.rels.values():
        if 'image' in rel.reltype:
            image_data = rel.target_part.blob
            ext = 'png'
            if image_data[:2] == b'\xff\xd8':
                ext = 'jpg'
            elif image_data[:4] == b'RIFF':
                ext = 'webp'
            out_path = os.path.join(out_dir, f'{Path(path).stem}_{count:02d}.{ext}')
            with open(out_path, 'wb') as f:
                f.write(image_data)
            count += 1
    print(f'Extracted {count} images from {path}')

for fname in os.listdir(LAB_DIR):
    fpath = os.path.join(LAB_DIR, fname)
    if fname.lower().endswith('.doc') and not fname.lower().endswith('.docx'):
        extract_images_doc(fpath, IMG_DIR)
    elif fname.lower().endswith('.docx'):
        extract_images_docx(fpath, IMG_DIR)

print('Done. Output:', OUT_DIR)
